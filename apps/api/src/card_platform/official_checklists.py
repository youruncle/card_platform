"""Deterministic official checklist facts with source locators and honest coverage."""

from collections import Counter
import csv
import io
from datetime import date
from hashlib import sha256
from pathlib import Path
import re
from urllib.parse import urlparse

PARSER_VERSION = "official-layout-v8-release-season"
ROOT = Path(__file__).resolve().parents[4]
STATUS_PATH = ROOT / "data/official-checklists/status.json"


def official_url(url: str) -> bool:
    p = urlparse(url)
    if p.scheme != "https" or p.username or p.password:
        return False
    return p.hostname in {
        "upperdeck.com",
        "www.upperdeck.com",
        "www.upperdeckepack.com",
        "topps.com",
        "www.topps.com",
        "paniniamerica.net",
        "www.paniniamerica.net",
        "blog.paniniamerica.net",
        "assets.paniniamerica.net",
        "support.paniniamerica.net",
    } or (p.hostname == "cdn.shopify.com" and p.path.startswith("/s/files/1/0662/9749/5709/files/"))


def window_status(release_date, start, end):
    if not release_date:
        return "unknown"
    return (
        "in_window"
        if date.fromisoformat(start) <= date.fromisoformat(release_date) <= date.fromisoformat(end)
        else "outside"
    )


def progress_summary(products, *, discovery_complete):
    counts = Counter(p.get("status", "pending") for p in products)
    total = len(products)
    return {
        "total": total,
        "eligible_total": total - counts["excluded"],
        "known_success_percent": round(100 * counts["published"] / (total - counts["excluded"]), 1) if total > counts["excluded"] else 0,
        "quarantined_rows": sum(p.get("issue_count", 0) for p in products),
        "counts": dict(counts),
        "discovery_complete": discovery_complete,
        "success_percent": round(100 * counts["published"] / total, 1)
        if total and discovery_complete
        else None,
        "processed_percent": round(
            100 * (counts["published"] + counts["failed"] + counts["excluded"]) / total, 1
        )
        if total and discovery_complete
        else None,
        "date_unknown": sum(p.get("window_status") == "unknown" for p in products),
    }


def parse_pages(pages):
    rows = []
    issues = []
    section = None
    pending = None
    for page_number, text in enumerate(pages, 1):
        for line_number, raw in enumerate(text.splitlines(), 1):
            line = raw.strip()
            if not line:
                continue
            if pending:
                prefix, prefix_page, prefix_line = pending
                if re.match(r"^[A-Z0-9]+\s{2,}", line):
                    line = prefix + line
                else:
                    issues.append(
                        {
                            "page": prefix_page,
                            "line": prefix_line,
                            "text": prefix,
                            "reason": "unresolved_wrapped_card_number",
                        }
                    )
                pending = None
            if re.fullmatch(r"[A-Z0-9]+(?:-[A-Z0-9]+)*-", line):
                pending = (line, page_number, line_number)
                continue
            if line in {"SELL SHEET", "SELL SHEETS"}:
                section = None
                continue
            if (
                line.isupper()
                and not re.search(r"\s{2,}", line)
                and not any(
                    token in line
                    for token in ["CARD NUMBER", "PLAYER NAME", "CHECKLIST", "TEAM NAME"]
                )
            ):
                section = line
                continue
            # Layout mode preserves field separation. A flattened line is not
            # safe to split at arbitrary name/team words.
            match = re.match(r"^([A-Za-z0-9]+(?:[-/][A-Za-z0-9]+)*)\s+(.+)$", line)
            if match and (re.search(r"\d", match[1]) or "-" in match[1]):
                if re.search(r"[A-Z][a-z]{2,}$", match[1]):
                    issues.append(
                        {
                            "page": page_number,
                            "line": line_number,
                            "text": line,
                            "reason": "glued_card_number_and_player",
                        }
                    )
                    continue
                columns = re.split(r"\s{3,}", match[2])
                if section and len(columns) >= 2 and bool(columns[0].strip()):
                    row = {
                        "card_number": match[1],
                        "player": columns[0].strip(),
                        "team": columns[1].strip(),
                        "section": section,
                        "insert": None
                        if section.startswith("BASE") and not re.search(r"AUTOGRAPH|RELIC", section)
                        else section.title(),
                        "source_page": page_number,
                        "source_line": line_number,
                        "raw_text": line,
                    }
                    rows.append(row)
                else:
                    issues.append(
                        {
                            "page": page_number,
                            "line": line_number,
                            "text": line,
                            "reason": "unresolved_columns_or_section",
                        }
                    )
            elif re.match(r"^(?:\d+|[A-Z]+-\w+)\s", line):
                issues.append(
                    {
                        "page": page_number,
                        "line": line_number,
                        "text": line,
                        "reason": "unresolved_card_row",
                    }
                )
            elif line.isupper() and re.search(r"[A-Z]", line) and len(line) < 120:
                if not any(
                    token in line
                    for token in [
                        "CARD NUMBER",
                        "PLAYER NAME",
                        "SUBJECT NAME",
                        "CHECKLIST",
                        "TEAM NAME",
                    ]
                ):
                    section = line
    if pending:
        issues.append(
            {
                "page": pending[1],
                "line": pending[2],
                "text": pending[0],
                "reason": "unresolved_wrapped_card_number",
            }
        )
    return rows, issues


def parse_document(path):
    path = Path(path)
    if path.read_bytes()[:4] == b"%PDF":
        from pypdf import PdfReader

        reader = PdfReader(path)
        pages = [p.extract_text(extraction_mode="layout") for p in reader.pages]
        rows, issues = parse_pages(pages)
        return rows, issues, pages
    if path.read_bytes()[:2] == b"PK":
        from openpyxl import load_workbook

        workbook = load_workbook(path, read_only=True, data_only=True)
        pages = []
        for sheet in workbook:
            lines = [sheet.title.upper()]
            for cells in sheet.iter_rows(values_only=True):
                values = [str(value).strip() for value in cells if value is not None]
                if values:
                    lines.append("    ".join(values))
            pages.append("\n".join(lines))
        workbook.close()
        rows, issues = parse_pages(pages)
        return rows, issues, pages
    text = path.read_text(encoding="utf-8-sig")
    if text.startswith(('"SPORT",', 'SPORT,')):
        return parse_panini_csv(text)
    if 'class="csv-table"' in text:
        return parse_upper_deck_html(text)
    raise ValueError("Unsupported official document format; raw file retained")


def parse_panini_csv(text):
    records = list(csv.DictReader(io.StringIO(text)))
    sections = {r.get("CARD SET", "").strip() for r in records}
    variants = {"Amber", "Artist Proof", "Canvas", "Gold", "Jade", "Masterpiece",
                "Obsidian", "Pink", "Red and Blue", "Ruby", "Sapphire", "Violet",
                "Silver", "Red", "Blue", "Green", "Black", "Purple", "Orange",
                "Plaid", "Cracked Ice"}
    rows, issues = [], []
    for line, record in enumerate(records, 2):
        section = record.get("CARD SET", "").strip()
        raw = str(record)
        # Some official exports include unrelated Instant sets inside a product.
        # Keep the entire CSV but do not assert these inconsistent identities.
        if re.match(r"^20\d{2}.*Basketball\s*-", section, re.I):
            issues.append({"page": 1, "line": line, "text": raw, "reason": "cross_product_section"})
            continue
        if record.get("SPORT") != "Basketball" or not all(record.get(k) for k in ("ATHLETE", "CARD NUMBER", "CARD SET")):
            issues.append({"page": 1, "line": line, "text": raw, "reason": "invalid_csv_identity"})
            continue
        family, parallel = section, None
        for suffix in sorted(variants, key=len, reverse=True):
            if section.endswith(" " + suffix) and section[:-(len(suffix) + 1)] in sections:
                family, parallel = section[:-(len(suffix) + 1)], suffix
                break
        sequence = record.get("SEQUENCE", "").strip()
        insert = family
        if family.casefold() == "base" or family.casefold().startswith("base "):
            insert = None
        elif parallel is None and any(section.startswith(parent + " ") for parent in sections if parent):
            # An unclassified suffix may be a parallel; never promote it to an insert.
            insert = None
        rows.append({"card_number": record["CARD NUMBER"].strip(), "player": record["ATHLETE"].strip(),
                     "team": record.get("TEAM", "").strip(), "section": section,
                     "insert": insert,
                     "parallel": parallel,
                     "serial_total": int(sequence) if sequence.isdigit() and int(sequence) > 0 else None,
                     "source_page": 1, "source_line": line, "raw_text": raw})
    return rows, issues, [text]


def product_fields(title):
    from card_platform.domain.season import explicit_seasons

    m = re.search(r"((?:19|20)\d{2})(?:[-/](\d{2,4}))?", title)
    start = int(m[1]) if m else None
    end = int(m[2]) if m and m[2] else start
    if end is not None and end < 100:
        end = (start // 100) * 100 + end
        if end < start:
            end += 100
    seasons = explicit_seasons(title)
    if len(seasons) > 1:
        raise ValueError('Conflicting explicit release seasons in official title')
    if seasons:
        season = next(iter(seasons))
        if start not in season or season[1] != season[0] + 1:
            raise ValueError('Official directory year conflicts with release season')
        start, end = season
    product = re.sub(r"^\s*\d{4}(?:[-/]\d{2,4})?\s*", "", title).replace("®", "")
    if seasons:
        product = re.sub(r"\(\s*(?:19|20)\d{2}\s*[-‐‑‒–—−/]\s*\d{2,4}\s*\)", "", product)
        product = re.sub(r"\s+", " ", product)
    product = re.sub(r"\s+Basketball\b", "", product, flags=re.I).strip()
    # Official directory typo; the linked document is titled Hoops.
    product = product.replace("Hopps", "Hoops")
    return {"product": product, "season_start": start, "season_end": end}


def build_manifest(product, rows):
    fields = product_fields(product["title"])
    source_key = product["key"]
    entities = {}

    def add(kind, name, attrs, claim):
        key = sha256((kind + name).encode()).hexdigest()[:24]
        entities[key] = {
            "key": key,
            "entity_type": kind,
            "canonical_name": name,
            "attributes": attrs,
            "source_key": source_key,
            "review_status": "approved",
            "confidence": 1,
            "claim": claim,
            "aliases": [],
        }
        return key

    add(
        "product",
        product["title"],
        {**fields, "manufacturer": product["publisher"], "sport": "basketball"},
        product["title"],
    )
    for row in rows:
        locator = f"page {row['source_page']} line {row['source_line']}"
        attrs = {
            **fields,
            "brand": product["publisher"],
            "sport": "basketball",
            "player": row["player"],
            "card_number": row["card_number"],
            "team": row["team"],
            "insert": row["insert"],
            "section": row["section"],
            "source_url": product["source_url"],
            "source_locator": locator,
            "source_checksum": product["checksum"],
            "parser_version": PARSER_VERSION,
            "authority": "official_checklist",
            "approval_method": "automatic_official_source",
            "raw_text": row["raw_text"],
        }
        if product["publisher"] == "Upper Deck":
            attrs["brand_aliases"] = ["UD"] + [b for b in ("Skybox", "Fleer") if re.search(r"\b" + b + r"\b", product["title"], re.I)]
        if row.get('sport_evidence'):
            attrs['sport_evidence'] = row['sport_evidence']
        for field in ("parallel", "serial_total"):
            if row.get(field) is not None:
                attrs[field] = row[field]
        name = f"{product['title']} · {row['section']} · #{row['card_number']} · {row['player']}"
        add("card", name, attrs, f"{locator}: {row['raw_text']}")
        if row["insert"]:
            add(
                "insert",
                f"{product['title']} · {row['insert']}",
                {
                    **fields,
                    "insert_name": row["insert"],
                    "sport": "basketball",
                    "source_url": product["source_url"],
                },
                f"Official section {row['section']}",
            )
    return {
        "batch_key": "official-" + source_key,
        "manifest_version": 1,
        "catalog_enabled": False,
        "required_regressions": [],
        "sources": [
            {
                "key": source_key,
                "title": product["title"],
                "url": product["source_url"],
                "publisher": product["publisher"],
                "source_type": "official_checklist",
                "checksum": product["checksum"],
                "retrieved_at": product.get("retrieved_at") or date.today().isoformat(),
                "metadata": {
                    "directory_url": product.get("directory_url"),
                    "parser_version": PARSER_VERSION,
                    "release_date": product.get("release_date"),
                    "date_basis": product.get("date_basis"),
                    "date_scope": product.get("date_scope"),
                    "ownership_evidence": product.get("ownership_evidence"),
                    "document_format": product.get("document_format"),
                    "retrieval_format": product.get("retrieval_format"),
                },
            }
        ],
        "entities": list(entities.values()),
        "relations": [],
        "coverage": [],
    }


def discover_topps_links(html: str, directory_url: str):
    """Parse public official links; season filtering is separate from release dates."""
    from html.parser import HTMLParser
    from urllib.parse import urljoin

    class Links(HTMLParser):
        def __init__(self):
            super().__init__()
            self.href = None
            self.label = []
            self.products = {}

        def handle_starttag(self, tag, attrs):
            if tag == "a":
                self.href = dict(attrs).get("href")
                self.label = []

        def handle_data(self, data):
            if self.href:
                self.label.append(data)

        def handle_endtag(self, tag):
            if tag != "a" or not self.href:
                return
            url = urljoin(directory_url, self.href)
            title = " ".join(" ".join(self.label).split())
            if (
                "basketball" in title.casefold()
                and official_url(url)
                and re.search(r"\.(pdf|xlsx?)(?:\?|$)", url, re.I)
            ):
                key = "topps-" + sha256(url.split("?")[0].encode()).hexdigest()[:16]
                self.products[key] = {
                    "key": key,
                    "title": title,
                    "source_url": url,
                    "publisher": "Topps",
                    "directory_url": directory_url,
                    "release_date": None,
                    "window_status": "unknown",
                    "status": "pending",
                }
            self.href = None

    parser = Links()
    parser.feed(html)
    return list(parser.products.values())


def isolate_conflicting_rows(rows):
    """Do not let manifest key deduplication silently choose conflicting facts."""
    grouped = {}
    for row in rows:
        key = tuple(row.get(k) for k in ('section', 'card_number', 'player'))
        grouped.setdefault(key, []).append(row)
    accepted, issues = [], []
    for group in grouped.values():
        facts = {tuple(row.get(k) for k in ('team', 'insert', 'parallel', 'serial_total')) for row in group}
        if len(facts) > 1:
            issues.extend(dict(page=r['source_page'], line=r['source_line'], text=r['raw_text'], reason='conflicting_card_identity') for r in group)
        else:
            accepted.extend(group)
    return accepted, issues


def parse_teamless_pages(pages):
    """Explicit two-column Topps layout; only used for verified teamless sources."""
    rows, issues = [], []
    section = category = None
    for page, text in enumerate(pages, 1):
        for line, raw in enumerate(text.splitlines(), 1):
            value = raw.strip()
            if not value: continue
            if value in {'BASE', 'INSERT', 'AUTOGRAPH'}:
                category = value; section = value; continue
            if value == 'CARDS': continue
            if value.isupper() and re.search('[A-Z]',value):
                section=value;continue
            match = re.fullmatch(r'([A-Za-z0-9]+(?:[-/][A-Za-z0-9]+)*)\s+([A-Za-zÀ-ž][A-Za-zÀ-ž .’\'-]+)',value)
            if match and section:
                rows.append(dict(card_number=match[1],player=match[2].strip(),team=None,section=section,
                    insert=None if category=='BASE' else section.title(),source_page=page,source_line=line,raw_text=raw))
            else: issues.append(dict(page=page,line=line,text=raw,reason='unresolved_teamless_row'))
    return rows,issues


def parse_mercury_pages(pages):
    """Mercury 2023-24's explicitly labelled card tables; marketing stays raw."""
    rows,issues=[],[]
    section=None
    headers={'Base Cards','Insert Cards Base /99 Refractor','Autograph Cards','Dual Autograph Cards','Autograph-Relic Cards'}
    for page,text in enumerate(pages,1):
        for line,raw in enumerate(text.splitlines(),1):
            value=raw.strip()
            label=' '.join(value.split())
            if label in headers:
                section=label;continue
            if 'Card Parallels' in label:
                section=None;continue
            if not section or not value: continue
            match=re.match(r'^([0-9]+|[A-Z]+-[A-Z0-9]+)\s+(.+)$',value)
            if not match:
                issues.append(dict(page=page,line=line,text=raw,reason='unresolved_mercury_row'));continue
            columns=re.split(r'\s{3,}',match[2])
            player=re.sub(r'\s+RC$', '',columns[0]).strip()
            insert=None if section=='Base Cards' else section
            if section.startswith('Insert Cards'):
                if len(columns)!=2 or not columns[1].endswith('Insert'):
                    issues.append(dict(page=page,line=line,text=raw,reason='unresolved_mercury_insert'));continue
                insert=columns[1][:-6].strip()
            rows.append(dict(card_number=match[1],player=player,team=None,section=insert or section,insert=insert,
                source_page=page,source_line=line,raw_text=raw))
    return rows,issues


def parse_upper_deck_html(text, *, row_sport_evidence=()):
    """Read explicit official table columns; never infer sport or update scope."""
    from html import unescape
    if '[Truncated]' in text or text.count('<table') != text.count('</table>'):
        raise ValueError('Incomplete official HTML table archive')
    rows, issues, pages = [], [], []
    proofs = {(p.get('section'),p.get('card_number'),p.get('player')):p for p in row_sport_evidence
              if p.get('sport') == 'basketball' and official_url(p.get('source_url',''))
              and re.fullmatch(r'[a-f0-9]{64}',p.get('source_checksum','')) and p.get('source_locator') and p.get('raw_text')}
    for page, table in enumerate(re.findall(r'<table\b[^>]*class="csv-table"[^>]*>(.*?)</table>', text, re.S), 1):
        records = [[unescape(re.sub(r'<[^>]+>', '', c)).strip().lstrip('\ufeff')
                    for c in re.findall(r'<t[dh]\b[^>]*>(.*?)</t[dh]>', row, re.S)]
                   for row in re.findall(r'<tr\b[^>]*>(.*?)</tr>', table, re.S)]
        if not records:
            continue
        headers = [h.casefold() for h in records[0]]
        pages.append('\n'.join('\t'.join(r) for r in records))
        required = {'set name', 'card', 'description'}
        if not required.issubset(headers):
            issues.append(dict(reason='unsupported_columns',source_page=page,raw_text=records[0]));continue
        for line, cells in enumerate(records[1:], 2):
            evidence = dict(source_page=page,source_line=line,raw_text='\t'.join(cells))
            if len(cells) != len(headers):
                issues.append(dict(reason='column_mismatch',**evidence));continue
            record = dict(zip(headers,cells))
            section, player, number = (record[k] for k in ('set name','description','card'))
            sport = record.get('category','').casefold()
            proof = proofs.get((section,number,player)) if not sport else None
            if proof:
                sport = 'basketball'
                evidence['sport_evidence'] = proof
            if sport and sport != 'basketball':
                continue
            reason = None
            if not sport: reason='sport_unknown'
            elif re.search(r'\b(?:19|20)\d{2}\b.*\bupdate\b',section,re.I): reason='cross_product_update'
            elif not all((section,player,number)) or number.casefold() in {'nno','n/a','-'}: reason='missing_identity'
            if reason:
                issues.append(dict(reason=reason,**evidence));continue
            base = re.match(r'^base(?:\s+set)?\b[\s-]*(.*)',section,re.I)
            # A base suffix can name a subset (Fresh Foundation, First Look,
            # Short Prints). Preserve it as section, without inventing a parallel.
            parallel = None
            # Named non-base parallels preserve the full section but do not invent
            # an insert identity when the table has no separate insert column.
            ambiguous_parallel = bool(re.search(r'\bparallel\b',section,re.I))
            serial = record.get("#'d", record.get("serial #'d", ''))
            rows.append(dict(section=section,card_number=number,player=player,team=None,
                insert=None if base or ambiguous_parallel else section,
                parallel=parallel,serial_total=int(serial) if serial.isdigit() else None,**evidence))
    return rows,issues,pages


def parse_upper_deck_epack(document, product_title):
    """e-Pack's official Basketball directory supplies sport; rows prove product."""
    import json
    rows, issues, pages = [], [], []
    for page, (group, records) in enumerate(document.items(), 1):
        pages.append(json.dumps({group: records},ensure_ascii=False))
        if not isinstance(records,list):
            raise ValueError('Invalid e-Pack checklist group')
        for line, record in enumerate(records, 1):
            evidence=dict(source_page=page,source_line=line,raw_text=json.dumps(record,ensure_ascii=False,sort_keys=True))
            if record.get('EpackProductName') != product_title or record.get('InsertName') != group:
                issues.append(dict(reason='cross_product_scope',**evidence));continue
            section=group.removeprefix(product_title).strip(' -')
            player,number=record.get('PlayerName'),record.get('CardNumber')
            if not all((player,number,section)):
                issues.append(dict(reason='missing_identity',**evidence));continue
            base=re.match(r'^Base(?: Set)?\b[\s-]*(.*)',section,re.I)
            parallel=None
            if 'Parallel' in section:
                parallel=re.sub(r'\s*Parallel\b','',base[1] if base else section).strip(' -') or None
            rows.append(dict(section=section,card_number=number,player=player,team=None,
                insert=None if base or parallel else section,parallel=parallel,serial_total=None,**evidence))
    return rows,issues,pages


def publisher_progress(products):
    result = {}
    for publisher in ('Topps', 'Panini', 'Upper Deck'):
        selected = [p for p in products if p.get('publisher') == publisher]
        summary = progress_summary(selected, discovery_complete=False)
        years = {}
        for p in selected:
            year = str(product_fields(p['title'])['season_start'] or 'unknown')
            counts = years.setdefault(year, {'discovered': 0})
            counts['discovered'] += 1
            status = p.get('status', 'pending')
            counts[status] = counts.get(status, 0) + 1
        result[publisher] = dict(summary, years=years, full_window_percent=None)
    return result
